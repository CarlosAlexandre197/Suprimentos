"""
Controle de Suprimentos
- Cadastro de transferências (Produto, Quantidade, Solicitante, Destino, Despachado, Data)
- Banco de dados SQLite (arquivo suprimentos.db, criado ao lado deste script)
- Relatório mensal com totais por produto e por destino, exportável para Excel, PDF ou CSV
- Importação das abas "Transferências ..." da planilha antiga

Requisitos: Python 3.8+ (tkinter e sqlite3 já vêm com o Python).
Opcional: pip install openpyxl reportlab   (openpyxl = Excel, reportlab = PDF)
"""
import csv
import os
import sqlite3
import sys
import tkinter as tk
from datetime import date, datetime
from tkinter import filedialog, messagebox, ttk

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import reportlab  # usado para gerar o PDF
except ImportError:
    reportlab = None

BASE_DIR = os.path.dirname(sys.executable) if getattr(sys, "frozen", False) else os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(BASE_DIR, "suprimentos.db")
# paleta azul
AZUL = "#1F5FBF"
AZUL_ESCURO = "#12397A"
AZUL_HOVER = "#3B7BDB"
AZUL_CLARO = "#DCE8F8"
AZUL_CLARO_BORDA = "#1F5FBF"
LINHA_ALT = "#EEF4FC"
FUNDO = "#F5F8FD"
TEXTO = "#1B2A41"

MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
         "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]


# ---------------------------------------------------------------- banco de dados
def consultar(sql, params=()):
    con = sqlite3.connect(DB)
    try:
        return con.execute(sql, params).fetchall()
    finally:
        con.close()


def executar(sql, params=()):
    con = sqlite3.connect(DB)
    try:
        con.execute(sql, params)
        con.commit()
    finally:
        con.close()


def iniciar_db():
    executar("""
        CREATE TABLE IF NOT EXISTS transferencias (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            produto TEXT NOT NULL,
            quantidade REAL NOT NULL,
            solicitante TEXT,
            destino TEXT,
            despachado TEXT DEFAULT 'Sim',
            data TEXT NOT NULL  -- formato AAAA-MM-DD
        )""")


def fmt_qtd(q):
    return f"{q:g}".replace(".", ",")


def br_para_iso(texto):
    return datetime.strptime(texto.strip(), "%d/%m/%Y").strftime("%Y-%m-%d")


def iso_para_br(texto):
    return datetime.strptime(texto, "%Y-%m-%d").strftime("%d/%m/%Y")


# ---------------------------------------------------------------- interface
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Controle de Suprimentos")
        self.geometry("1050x700")
        self.minsize(900, 600)
        self.aplicar_estilo()

        # barra azul no topo
        topo_bar = tk.Frame(self, bg=AZUL_ESCURO, height=64)
        topo_bar.pack(fill="x")
        topo_bar.pack_propagate(False)
        tk.Label(topo_bar, text="Controle de Suprimentos", bg=AZUL_ESCURO, fg="white",
                 font=("Segoe UI", 18, "bold")).pack(side="left", padx=20)
        tk.Label(topo_bar, text="Transferências e relatórios mensais", bg=AZUL_ESCURO,
                 fg="#BBD3F5", font=("Segoe UI", 10)).pack(side="left", padx=4, pady=(10, 0))
        tk.Frame(self, bg=AZUL_CLARO_BORDA, height=3).pack(fill="x")

        abas = ttk.Notebook(self)
        abas.pack(fill="both", expand=True, padx=12, pady=12)
        self.aba_cad = ttk.Frame(abas)
        self.aba_rel = ttk.Frame(abas)
        abas.add(self.aba_cad, text="Cadastro")
        abas.add(self.aba_rel, text="Relatório mensal")
        self.montar_cadastro()
        self.montar_relatorio()
        self.atualizar_listas()
        self.carregar_lista()

    # ------------------------------------------------------------ estilo
    def aplicar_estilo(self):
        s = ttk.Style(self)
        s.theme_use("clam")
        self.configure(bg=FUNDO)
        fonte = ("Segoe UI", 10)
        s.configure(".", font=fonte, background=FUNDO, foreground=TEXTO)
        s.configure("TFrame", background=FUNDO)
        s.configure("TLabel", background=FUNDO, foreground=TEXTO)

        # abas
        s.configure("TNotebook", background=FUNDO, borderwidth=0, tabmargins=(0, 0, 0, 0))
        s.configure("TNotebook.Tab", background=AZUL_CLARO, foreground=AZUL_ESCURO,
                    padding=(22, 9), font=("Segoe UI", 10, "bold"), borderwidth=0)
        s.map("TNotebook.Tab",
              background=[("selected", AZUL), ("active", "#B9D3F3")],
              foreground=[("selected", "white")])

        # quadro do formulário
        s.configure("TLabelframe", background=FUNDO, bordercolor=AZUL, relief="solid", borderwidth=1)
        s.configure("TLabelframe.Label", background=FUNDO, foreground=AZUL_ESCURO,
                    font=("Segoe UI", 11, "bold"))

        # botões azuis
        s.configure("TButton", background=AZUL, foreground="white", padding=(14, 6),
                    font=("Segoe UI", 10, "bold"), borderwidth=0, focusthickness=0)
        s.map("TButton",
              background=[("pressed", AZUL_ESCURO), ("active", AZUL_HOVER)],
              foreground=[("disabled", "#E0E0E0")])

        # campos
        for w in ("TEntry", "TCombobox", "TSpinbox"):
            s.configure(w, fieldbackground="white", bordercolor="#9DB7DC",
                        lightcolor="#9DB7DC", darkcolor="#9DB7DC", padding=4)
        s.map("TCombobox", bordercolor=[("focus", AZUL)], lightcolor=[("focus", AZUL)],
              darkcolor=[("focus", AZUL)])
        s.map("TEntry", bordercolor=[("focus", AZUL)], lightcolor=[("focus", AZUL)],
              darkcolor=[("focus", AZUL)])

        # tabelas: cabeçalho azul
        s.configure("Treeview", background="white", fieldbackground="white",
                    foreground=TEXTO, rowheight=26, borderwidth=0)
        s.configure("Treeview.Heading", background=AZUL, foreground="white",
                    font=("Segoe UI", 10, "bold"), padding=(6, 7), relief="flat")
        s.map("Treeview.Heading", background=[("active", AZUL_HOVER)])
        s.map("Treeview", background=[("selected", "#2F6FD0")],
              foreground=[("selected", "white")])

        # barras de rolagem azuis
        s.configure("Vertical.TScrollbar", background=AZUL, troughcolor=AZUL_CLARO,
                    bordercolor=AZUL_CLARO, arrowcolor="white", relief="flat")
        s.map("Vertical.TScrollbar", background=[("active", AZUL_HOVER)])

    def criar_tabela(self, pai, **kw):
        """Treeview com barra de rolagem azul e linhas zebradas."""
        quadro = ttk.Frame(pai)
        tree = ttk.Treeview(quadro, **kw)
        barra = ttk.Scrollbar(quadro, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=barra.set)
        tree.pack(side="left", fill="both", expand=True)
        barra.pack(side="right", fill="y")
        tree.tag_configure("par", background="white")
        tree.tag_configure("impar", background=LINHA_ALT)
        return quadro, tree

    @staticmethod
    def inserir(tree, valores):
        n = len(tree.get_children())
        tree.insert("", "end", values=valores, tags=("par" if n % 2 == 0 else "impar",))

    # ------------------------------------------------------------ aba cadastro
    def montar_cadastro(self):
        f = ttk.LabelFrame(self.aba_cad, text="Nova transferência")
        f.pack(fill="x", padx=8, pady=8)
        f.columnconfigure(1, weight=1)
        f.columnconfigure(3, weight=1)

        self.v_produto = tk.StringVar()
        self.v_qtd = tk.StringVar(value="1")
        self.v_solic = tk.StringVar()
        self.v_destino = tk.StringVar()
        self.v_desp = tk.StringVar(value="Sim")
        self.v_data = tk.StringVar(value=date.today().strftime("%d/%m/%Y"))

        ttk.Label(f, text="Produto:").grid(row=0, column=0, sticky="e", padx=4, pady=4)
        self.cb_produto = ttk.Combobox(f, textvariable=self.v_produto)
        self.cb_produto.grid(row=0, column=1, sticky="ew", padx=4, pady=4)

        ttk.Label(f, text="Quantidade:").grid(row=0, column=2, sticky="e", padx=4, pady=4)
        ttk.Entry(f, textvariable=self.v_qtd, width=10).grid(row=0, column=3, sticky="w", padx=4, pady=4)

        ttk.Label(f, text="Solicitante:").grid(row=1, column=0, sticky="e", padx=4, pady=4)
        self.cb_solic = ttk.Combobox(f, textvariable=self.v_solic)
        self.cb_solic.grid(row=1, column=1, sticky="ew", padx=4, pady=4)

        ttk.Label(f, text="Destino:").grid(row=1, column=2, sticky="e", padx=4, pady=4)
        self.cb_destino = ttk.Combobox(f, textvariable=self.v_destino)
        self.cb_destino.grid(row=1, column=3, sticky="ew", padx=4, pady=4)

        ttk.Label(f, text="Despachado:").grid(row=2, column=0, sticky="e", padx=4, pady=4)
        ttk.Combobox(f, textvariable=self.v_desp, values=["Sim", "Não"], state="readonly",
                     width=8).grid(row=2, column=1, sticky="w", padx=4, pady=4)

        ttk.Label(f, text="Data (dd/mm/aaaa):").grid(row=2, column=2, sticky="e", padx=4, pady=4)
        ttk.Entry(f, textvariable=self.v_data, width=12).grid(row=2, column=3, sticky="w", padx=4, pady=4)

        botoes = ttk.Frame(f)
        botoes.grid(row=3, column=0, columnspan=4, pady=6)
        ttk.Button(botoes, text="Salvar", command=self.salvar).pack(side="left", padx=4)
        ttk.Button(botoes, text="Limpar", command=self.limpar).pack(side="left", padx=4)
        ttk.Button(botoes, text="Importar da planilha Excel...", command=self.importar_excel).pack(side="left", padx=16)
        self.bind("<Return>", lambda e: self.salvar() if self.focus_get() is not None else None)

        colunas = ("id", "data", "produto", "quantidade", "solicitante", "destino", "despachado")
        titulos = ("ID", "Data", "Produto", "Qtd", "Solicitante", "Destino", "Despachado")
        larguras = (50, 90, 260, 70, 100, 260, 90)
        quadro, self.tree = self.criar_tabela(self.aba_cad, columns=colunas, show="headings", height=14)
        for c, t, w in zip(colunas, titulos, larguras):
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w, anchor="w" if c in ("produto", "destino") else "center")
        quadro.pack(fill="both", expand=True, padx=8)
        rodape = ttk.Frame(self.aba_cad)
        rodape.pack(fill="x", padx=8, pady=6)
        ttk.Label(rodape, text="Últimos 300 registros").pack(side="left")
        ttk.Button(rodape, text="Excluir selecionado", command=self.excluir).pack(side="right")

    def atualizar_listas(self):
        def valores(coluna):
            return [r[0] for r in consultar(
                f"SELECT DISTINCT {coluna} FROM transferencias "
                f"WHERE {coluna} IS NOT NULL AND {coluna} != '' ORDER BY {coluna}")]
        self.cb_produto["values"] = valores("produto")
        self.cb_solic["values"] = valores("solicitante")
        self.cb_destino["values"] = valores("destino")

    def carregar_lista(self):
        self.tree.delete(*self.tree.get_children())
        for r in consultar("SELECT id, data, produto, quantidade, solicitante, destino, despachado "
                           "FROM transferencias ORDER BY data DESC, id DESC LIMIT 300"):
            self.inserir(self.tree, (r[0], iso_para_br(r[1]), r[2], fmt_qtd(r[3]), r[4], r[5], r[6]))

    def salvar(self):
        produto = self.v_produto.get().strip()
        if not produto:
            messagebox.showwarning("Atenção", "Informe o produto.")
            return
        try:
            qtd = float(self.v_qtd.get().replace(",", "."))
            if qtd <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Atenção", "Quantidade inválida.")
            return
        try:
            data_iso = br_para_iso(self.v_data.get())
        except ValueError:
            messagebox.showwarning("Atenção", "Data inválida. Use dd/mm/aaaa.")
            return
        executar("INSERT INTO transferencias (produto, quantidade, solicitante, destino, despachado, data) "
                 "VALUES (?,?,?,?,?,?)",
                 (produto, qtd, self.v_solic.get().strip(), self.v_destino.get().strip(),
                  self.v_desp.get(), data_iso))
        # mantém solicitante/destino/data para lançar vários itens seguidos
        self.v_produto.set("")
        self.v_qtd.set("1")
        self.atualizar_listas()
        self.carregar_lista()
        self.cb_produto.focus_set()

    def limpar(self):
        self.v_produto.set("")
        self.v_qtd.set("1")
        self.v_solic.set("")
        self.v_destino.set("")
        self.v_desp.set("Sim")
        self.v_data.set(date.today().strftime("%d/%m/%Y"))

    def excluir(self):
        sel = self.tree.selection()
        if not sel:
            return
        if messagebox.askyesno("Confirmar", f"Excluir {len(sel)} registro(s)?"):
            for item in sel:
                executar("DELETE FROM transferencias WHERE id = ?", (self.tree.item(item, "values")[0],))
            self.atualizar_listas()
            self.carregar_lista()

    def importar_excel(self):
        if openpyxl is None:
            self.avisar_instalacao("openpyxl", "importar Excel")
            return
        caminho = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx *.xlsm")])
        if not caminho:
            return
        wb = openpyxl.load_workbook(caminho, data_only=True)
        novos = 0
        for ws in wb.worksheets:
            if not str(ws["A2"].value or "").strip().lower().startswith("produto"):
                continue  # ignora abas fora do padrão (ex.: aba "Produtos")
            if not str(ws["E2"].value or "").strip().lower().startswith("despach"):
                continue
            for linha in ws.iter_rows(min_row=3, max_col=6, values_only=True):
                produto, qtd, solic, destino, desp, data = linha
                if not produto or qtd in (None, ""):
                    continue
                try:
                    if isinstance(data, datetime):
                        data_iso = data.strftime("%Y-%m-%d")
                    elif isinstance(data, date):
                        data_iso = data.isoformat()
                    else:
                        data_iso = br_para_iso(str(data))
                    qtd = float(str(qtd).replace(",", "."))
                except (ValueError, TypeError):
                    continue
                params = (str(produto).strip(), qtd, solic, destino, desp or "Sim", data_iso)
                if consultar("SELECT 1 FROM transferencias WHERE produto=? AND quantidade=? AND "
                             "solicitante IS ? AND destino IS ? AND despachado=? AND data=?", params):
                    continue  # já importado antes
                executar("INSERT INTO transferencias (produto, quantidade, solicitante, destino, "
                         "despachado, data) VALUES (?,?,?,?,?,?)", params)
                novos += 1
        self.atualizar_listas()
        self.carregar_lista()
        messagebox.showinfo("Importação", f"{novos} registro(s) importado(s).")

    # ------------------------------------------------------------ aba relatório
    def montar_relatorio(self):
        topo = ttk.Frame(self.aba_rel)
        topo.pack(fill="x", padx=8, pady=8)
        hoje = date.today()
        self.v_mes = tk.StringVar(value=MESES[hoje.month - 1])
        self.v_ano = tk.StringVar(value=str(hoje.year))
        ttk.Label(topo, text="Mês:").pack(side="left")
        ttk.Combobox(topo, textvariable=self.v_mes, values=MESES, state="readonly",
                     width=12).pack(side="left", padx=4)
        ttk.Label(topo, text="Ano:").pack(side="left", padx=(8, 0))
        ttk.Spinbox(topo, from_=2020, to=2100, textvariable=self.v_ano, width=6).pack(side="left", padx=4)
        ttk.Button(topo, text="Gerar relatório", command=self.gerar_relatorio).pack(side="left", padx=8)
        ttk.Button(topo, text="Exportar (Excel/PDF/CSV)...", command=self.exportar).pack(side="left")
        self.lbl_resumo = ttk.Label(topo, text="", foreground=AZUL_ESCURO,
                                    font=("Segoe UI", 10, "bold"))
        self.lbl_resumo.pack(side="right")

        meio = ttk.Frame(self.aba_rel)
        meio.pack(fill="both", expand=True, padx=8)
        meio.columnconfigure(0, weight=1)
        meio.columnconfigure(1, weight=1)
        meio.rowconfigure(1, weight=1)

        ttk.Label(meio, text="Total por produto").grid(row=0, column=0, sticky="w")
        ttk.Label(meio, text="Total por destino").grid(row=0, column=1, sticky="w")
        q1, self.t_prod = self.criar_tabela(meio, columns=("p", "q"), show="headings")
        self.t_prod.heading("p", text="Produto")
        self.t_prod.heading("q", text="Qtd total")
        self.t_prod.column("p", width=250)
        self.t_prod.column("q", width=80, anchor="center")
        q1.grid(row=1, column=0, sticky="nsew", padx=(0, 4))
        q2, self.t_dest = self.criar_tabela(meio, columns=("d", "n", "q"), show="headings")
        self.t_dest.heading("d", text="Destino")
        self.t_dest.heading("n", text="Saídas")
        self.t_dest.heading("q", text="Qtd total")
        self.t_dest.column("d", width=250)
        self.t_dest.column("n", width=60, anchor="center")
        self.t_dest.column("q", width=80, anchor="center")
        q2.grid(row=1, column=1, sticky="nsew", padx=(4, 0))
        self.dados_rel = None

    def _mes_ano(self):
        return MESES.index(self.v_mes.get()) + 1, int(self.v_ano.get())

    def _consultas(self):
        mes, ano = self._mes_ano()
        ym = f"{ano:04d}-{mes:02d}"
        detalhe = consultar("SELECT data, produto, quantidade, solicitante, destino, despachado "
                            "FROM transferencias WHERE substr(data,1,7)=? ORDER BY data, id", (ym,))
        por_prod = consultar("SELECT produto, SUM(quantidade) FROM transferencias WHERE substr(data,1,7)=? "
                             "GROUP BY produto ORDER BY 2 DESC", (ym,))
        por_dest = consultar("SELECT destino, COUNT(*), SUM(quantidade) FROM transferencias "
                             "WHERE substr(data,1,7)=? GROUP BY destino ORDER BY 3 DESC", (ym,))
        return detalhe, por_prod, por_dest

    def gerar_relatorio(self):
        try:
            detalhe, por_prod, por_dest = self._consultas()
        except ValueError:
            messagebox.showwarning("Atenção", "Ano inválido.")
            return
        self.dados_rel = (detalhe, por_prod, por_dest)
        self.t_prod.delete(*self.t_prod.get_children())
        self.t_dest.delete(*self.t_dest.get_children())
        for p, q in por_prod:
            self.inserir(self.t_prod, (p, fmt_qtd(q)))
        for d, n, q in por_dest:
            self.inserir(self.t_dest, (d or "(sem destino)", n, fmt_qtd(q)))
        self.lbl_resumo.config(text=f"{len(detalhe)} saída(s) | {len(por_prod)} produto(s)")

    def avisar_instalacao(self, pacote, para):
        messagebox.showerror(
            f"Falta instalar o {pacote}",
            f"Para {para}, instale o {pacote} no mesmo Python que está rodando este programa.\n\n"
            f"Cole isto no terminal (Prompt de Comando):\n\n"
            f'"{sys.executable}" -m pip install {pacote}\n\n'
            "Depois feche e abra o programa novamente.")

    def exportar(self):
        self.gerar_relatorio()
        if not self.dados_rel or not self.dados_rel[0]:
            messagebox.showinfo("Relatório", "Não há registros neste mês.")
            return
        detalhe, por_prod, por_dest = self.dados_rel
        mes, ano = self._mes_ano()
        nome = f"Relatorio_Suprimentos_{ano}-{mes:02d}"
        # PDF, Excel e CSV sempre aparecem na janela; se faltar a biblioteca,
        # o programa avisa e explica como instalar.
        tipos = [("PDF", "*.pdf"), ("Excel", "*.xlsx"), ("CSV", "*.csv")]
        caminho = filedialog.asksaveasfilename(initialfile=nome, defaultextension=".pdf",
                                               filetypes=tipos)
        if not caminho:
            return
        cab = ["Data", "Produto", "Quantidade", "Solicitante", "Destino", "Despachado"]
        linhas = [(iso_para_br(r[0]), r[1], r[2], r[3], r[4], r[5]) for r in detalhe]
        if caminho.lower().endswith(".pdf"):
            if not reportlab:
                self.avisar_instalacao("reportlab", "gerar PDF")
                return
            self.exportar_pdf(caminho, mes, ano, detalhe, por_prod, por_dest)
        elif caminho.lower().endswith(".xlsx"):
            if not openpyxl:
                self.avisar_instalacao("openpyxl", "gerar Excel")
                return
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Detalhado"
            ws.append(cab)
            for l in linhas:
                ws.append(l)
            ws2 = wb.create_sheet("Por produto")
            ws2.append(["Produto", "Quantidade total"])
            for r in por_prod:
                ws2.append(r)
            ws3 = wb.create_sheet("Por destino")
            ws3.append(["Destino", "Saídas", "Quantidade total"])
            for r in por_dest:
                ws3.append(r)
            for w in (ws, ws2, ws3):
                for col in w.columns:
                    w.column_dimensions[col[0].column_letter].width = max(
                        12, min(50, max(len(str(c.value or "")) for c in col) + 2))
            wb.save(caminho)
        else:
            with open(caminho, "w", newline="", encoding="utf-8-sig") as fh:
                w = csv.writer(fh, delimiter=";")
                w.writerow(cab)
                w.writerows(linhas)
        messagebox.showinfo("Exportado", f"Relatório salvo em:\n{caminho}")

    def exportar_pdf(self, caminho, mes, ano, detalhe, por_prod, por_dest):
        from xml.sax.saxutils import escape

        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        estilos = getSampleStyleSheet()
        celula = ParagraphStyle("celula", parent=estilos["BodyText"], fontSize=9, leading=11)
        cab_est = ParagraphStyle("cab", parent=celula, fontName="Helvetica-Bold")

        def tabela(cabecalho, linhas, larguras):
            dados = [[Paragraph(escape(c), cab_est) for c in cabecalho]]
            for l in linhas:
                dados.append([Paragraph(escape("" if c is None else str(c)), celula) for c in l])
            t = Table(dados, colWidths=[w * cm for w in larguras], repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FFC000")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            return t

        def rodape(canvas, doc):
            canvas.setFont("Helvetica", 8)
            canvas.drawRightString(landscape(A4)[0] - 1.5 * cm, 1 * cm, f"Página {doc.page}")

        doc = SimpleDocTemplate(caminho, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                                topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                                title=f"Suprimentos - {MESES[mes - 1]}/{ano}")
        h1, h2 = estilos["Title"], estilos["Heading2"]
        historia = [
            Paragraph(f"Relatório de Suprimentos - {MESES[mes - 1]}/{ano}", h1),
            Paragraph(f"{len(detalhe)} saída(s) | {len(por_prod)} produto(s) | "
                      f"gerado em {date.today().strftime('%d/%m/%Y')}", estilos["Normal"]),
            Spacer(1, 0.5 * cm),
            Paragraph("Total por produto", h2),
            tabela(["Produto", "Quantidade total"], [(p, fmt_qtd(q)) for p, q in por_prod], [16, 5]),
            Spacer(1, 0.4 * cm),
            Paragraph("Total por destino", h2),
            tabela(["Destino", "Saídas", "Quantidade total"],
                   [(d or "(sem destino)", n, fmt_qtd(q)) for d, n, q in por_dest], [14, 3, 5]),
            Spacer(1, 0.4 * cm),
            Paragraph("Detalhamento", h2),
            tabela(["Data", "Produto", "Qtd", "Solicitante", "Destino", "Despachado"],
                   [(iso_para_br(r[0]), r[1], fmt_qtd(r[2]), r[3], r[4], r[5]) for r in detalhe],
                   [2.5, 9, 2.2, 3, 7, 3]),
        ]
        doc.build(historia, onFirstPage=rodape, onLaterPages=rodape)


if __name__ == "__main__":
    iniciar_db()
    App().mainloop()
